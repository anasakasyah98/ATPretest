##
# @file main.py
# @brief LTE AT command pre-test automation tool.
#
# Executes a predefined list of LTE AT command test cases against a modem
# through a serial port. The application records the modem responses,
# displays the test progress on the console, and generates a timestamped
# log file containing the test results.
#
# Features:
# - Executes predefined LTE AT commands
# - Logs all modem responses
# - Generates timestamped log files
# - Supports automatic modem reset handling (AT+CFUN=1,1)
#
# @author Anas Akasyah Bin Abdul Patas
# @version 1.0.0
# @date 2026-10-26
# @copyright Eagle Wireless
##

# Libraries
import sys
from pathlib import Path
from openpyxl import load_workbook
from test_case import at_cmd_general_info_list
import qcx216
import mdm9x07
import sdx35

general_info = at_cmd_general_info_list[0]

# Function

##
# @brief Read the jenkins value from the excel sheet.
#
# @param [in] excel_path Excel path
# @param [in] parameter_name Parameter name
#
# @return Jenkins parameter value
##
def read_jenkins_value(excel_path, parameter_name):

    workbook = load_workbook(excel_path, data_only=True)
    sheet = workbook.active

    header_row = 1

    # Find column indexes by header name
    headers = {
        cell.value: cell.column
        for cell in sheet[header_row]
        if cell.value is not None
    }

    param_col = headers["Jenkins Parameter"]
    value_col = headers["Value"]

    # Search parameter name
    for row in range(header_row + 1, sheet.max_row + 1):
        param = sheet.cell(row=row, column=param_col).value

        if param == parameter_name:
            jenkins_value = str(sheet.cell(row=row, column=value_col).value)
            workbook.close()
            return jenkins_value

    workbook.close()
    return "None"

##
# @brief Main application.
#
# @return None
##
def main_application():
    try:

        # Get username
        while True:
            username = input("Please Enter Username (-@eaglewireless.com) : ").strip().lower()

            if not username.endswith("@eaglewireless.com") or username == "@eaglewireless.com":
                print("\nWrong parameter. Try again\n")
            else:
                break

        port = input("Please Enter COM port (e.g. COM5)           : ")
        sim_ccid = input("Please Enter SIM CCID (20 digits)           : ")
        sim_imei = input("Please Enter SIM IMEI (15 digits)           : ")
        sim_imsi = input("Please Enter SIM IMSI (15 digits)           : ")

        # Detect platform
        while True:
            jira_id = input("Please Enter JIRA ID (e.g. SW9X07-10)       : ").strip().upper()

            if not any(platform in jira_id for platform in ("QCX216", "SW9X07", "SDX35")):
                print("\nWrong parameter. Try again\n")
            else:
                break

        # Run pre-test for the platform
        if "QCX216" in jira_id:
            print("\nQCX216 Pre-Test\n")
            qcx216.pretest_qcx216(username, port, sim_ccid, sim_imei, sim_imsi)
        elif "SW9X07" in jira_id:
            print("\nMDM9X07 Pre-Test\n")
            mdm9x07.pretest_mdm9x07(username, port, sim_ccid, sim_imei, sim_imsi)
        elif "SDX35" in jira_id:
            print("\nSDX35 Pre-Test\n")
            sdx35.pretest_sdx35(username, port, sim_ccid, sim_imei, sim_imsi)

    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    main_application()
    if getattr(sys, "frozen", False):
        input("\nPress Enter to exit...\n")
